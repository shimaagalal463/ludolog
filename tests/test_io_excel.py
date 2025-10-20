"""Tests for ludolog.io_excel."""
import pytest

openpyxl = pytest.importorskip("openpyxl")

from ludolog.errors import ImportError as LudologImportError
from ludolog.io_excel import export_collection_excel, export_plays_excel, sheet_names


def test_export_plays_excel_row_count(populated_library, tmp_path):
    out = tmp_path / "plays.xlsx"
    count = export_plays_excel(populated_library, out)
    assert count == 1
    assert out.exists()


def test_export_plays_excel_header_and_data(populated_library, tmp_path):
    out = tmp_path / "plays.xlsx"
    export_plays_excel(populated_library, out)
    from openpyxl import load_workbook

    wb = load_workbook(out)
    ws = wb.active
    assert ws.title == "Plays"
    headers = [cell.value for cell in ws[1]]
    assert headers == ["ID", "Date", "Game", "Players", "Winners", "Minutes", "Venue", "Notes"]
    assert ws.cell(2, 3).value == "Catan"
    assert "alice" in str(ws.cell(2, 4).value).lower()
    assert ws.cell(2, 6).value == 75
    assert ws.cell(2, 7).value == "Game Cafe"


def test_export_plays_excel_creates_parent_dirs(populated_library, tmp_path):
    out = tmp_path / "nested" / "dir" / "plays.xlsx"
    export_plays_excel(populated_library, out)
    assert out.exists()


def test_export_collection_excel_row_count(populated_library, tmp_path):
    out = tmp_path / "collection.xlsx"
    count = export_collection_excel(populated_library, out)
    assert count == 3


def test_export_collection_excel_includes_shelf(populated_library, tmp_path):
    out = tmp_path / "collection.xlsx"
    export_collection_excel(populated_library, out)
    from openpyxl import load_workbook

    wb = load_workbook(out)
    ws = wb.active
    assert ws.title == "Collection"
    headers = [cell.value for cell in ws[1]]
    assert "Shelf" in headers
    catan_row = next(
        row for row in ws.iter_rows(min_row=2, values_only=True) if row[0] == "Catan"
    )
    assert catan_row[6] == "Living Room"
    assert catan_row[7] == "yes"


def test_export_collection_excel_wishlist_not_owned(populated_library, tmp_path):
    out = tmp_path / "collection.xlsx"
    export_collection_excel(populated_library, out)
    from openpyxl import load_workbook

    wb = load_workbook(out)
    ws = wb.active
    ark_row = next(
        row for row in ws.iter_rows(min_row=2, values_only=True) if row[0] == "Ark Nova"
    )
    assert ark_row[7] == "no"


def test_sheet_names_reads_workbook(populated_library, tmp_path):
    out = tmp_path / "book.xlsx"
    export_plays_excel(populated_library, out)
    names = sheet_names(out)
    assert names == ["Plays"]


def test_sheet_names_missing_file_raises(tmp_path):
    with pytest.raises(FileNotFoundError):
        sheet_names(tmp_path / "nope.xlsx")


def test_export_plays_excel_empty_library(tmp_journal, tmp_path):
    out = tmp_path / "empty.xlsx"
    count = export_plays_excel(tmp_journal, out)
    assert count == 0
    from openpyxl import load_workbook

    ws = load_workbook(out).active
    assert ws.max_row == 1


def test_export_collection_excel_empty_library(tmp_journal, tmp_path):
    out = tmp_path / "empty_coll.xlsx"
    count = export_collection_excel(tmp_journal, out)
    assert count == 0


def test_export_plays_excel_without_openpyxl(monkeypatch, populated_library, tmp_path):
    import builtins

    real_import = builtins.__import__

    def fake_import(name, *args, **kwargs):
        if name == "openpyxl":
            raise ImportError("no openpyxl")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", fake_import)
    with pytest.raises(LudologImportError, match="openpyxl required"):
        export_plays_excel(populated_library, tmp_path / "x.xlsx")


def test_export_collection_creates_parent_dir(populated_library, tmp_path):
    out = tmp_path / "deep" / "coll.xlsx"
    export_collection_excel(populated_library, out)
    assert out.exists()
