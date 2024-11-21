"""Excel export via openpyxl."""
from __future__ import annotations

from pathlib import Path
from typing import TYPE_CHECKING

from .errors import ImportError as LudologImportError

if TYPE_CHECKING:
    from .storage import Library


def export_plays_excel(lib: "Library", path: Path | str) -> int:
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise LudologImportError("openpyxl required for excel export") from exc
    wb = Workbook()
    ws = wb.active
    ws.title = "Plays"
    ws.append(["ID", "Date", "Game", "Players", "Winners", "Minutes", "Venue", "Notes"])
    games = {g.id: g for g in lib.games}
    players = {p.id: p for p in lib.players}
    venues = {v.id: v for v in lib.venues}
    for play in lib.plays:
        g = games.get(play.game_id)
        pnames = ", ".join(players[i].display_name() for i in play.player_ids if i in players)
        wnames = ", ".join(players[i].display_name() for i in play.winner_ids if i in players)
        venue = venues.get(play.venue_id) if play.venue_id else None
        ws.append([
            play.id,
            play.played_date(),
            g.title if g else play.game_id,
            pnames,
            wnames,
            play.duration_minutes,
            venue.name if venue else "",
            play.notes,
        ])
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(p)
    return len(lib.plays)


def export_collection_excel(lib: "Library", path: Path | str) -> int:
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise LudologImportError("openpyxl required") from exc
    wb = Workbook()
    ws = wb.active
    ws.title = "Collection"
    ws.append(["Title", "BGG", "Designer", "Players", "Minutes", "Mechanics", "Shelf", "Owned"])
    shelves = {s.id: s for s in lib.shelves}
    for g in lib.games:
        shelf = shelves.get(g.shelf_id) if g.shelf_id else None
        ws.append([
            g.title,
            g.bgg_id,
            g.designer,
            g.player_range.label(),
            g.playtime.label(),
            ", ".join(g.mechanics),
            shelf.name if shelf else "",
            "yes" if g.owned else "no",
        ])
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return len(lib.games)


def sheet_names(path) -> list[str]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return []
    wb = load_workbook(path, read_only=True)
    return wb.sheetnames
